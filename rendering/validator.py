import docx
import os
from typing import Dict
from schema.document_spec import DocumentSpec

class PostRenderValidator:
    """
    Validates a generated DOCX file against the intended DocumentSpec.
    This acts as the final "Visual/Structural Validator" before the document is delivered to the user.
    """
    def __init__(self, docx_path: str, intended_spec: DocumentSpec):
        self.docx_path = docx_path
        self.intended_spec = intended_spec

    def validate(self) -> Dict[str, any]:
        """
        Runs structural checks on the generated DOCX.
        Returns a dict of validation results and any warnings.
        """
        if not os.path.exists(self.docx_path):
            return {"valid": False, "error": "DOCX file not found."}
            
        try:
            doc = docx.Document(self.docx_path)
        except Exception as e:
            return {"valid": False, "error": f"Failed to parse DOCX: {e}"}
            
        warnings = []
        
        # 1. Validate page layout (margins)
        if doc.sections:
            sec = doc.sections[0]
            # Convert python-docx EMU (English Metric Units) to mm
            top_margin_mm = sec.top_margin.mm if sec.top_margin else 0
            # Allow 2mm tolerance
            if abs(top_margin_mm - self.intended_spec.page.margins.top_mm) > 2.0:
                warnings.append(f"Top margin mismatch: found {top_margin_mm:.1f}mm, expected {self.intended_spec.page.margins.top_mm}mm")
                
        # 2. Validate basic structure (content exists)
        if len(doc.paragraphs) == 0 and len(doc.tables) == 0:
            warnings.append("Document appears empty (no paragraphs or tables).")
            
        # 3. Future Expansion: Agentic Visual QA
        # - Convert DOCX to PDF/Images
        # - Pass to a multi-modal model (like Qwen-VL) to check if the layout "looks" correct visually
        
        return {
            "valid": len(warnings) == 0,
            "warnings": warnings
        }

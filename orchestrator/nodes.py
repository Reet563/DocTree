import sys
import os
import json
from pydantic import ValidationError

sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from orchestrator.state import AgentState
from models.base import DocumentGenerator
from models.local.prompts import build_system_prompt
from schema.document_spec import DocumentSpec
from reference.vision_analyzer import VisionAnalyzer

class AgentNodes:
    def __init__(self, generator: DocumentGenerator):
        self.generator = generator

    def reference_agent_node(self, state: AgentState) -> AgentState:
        """
        Analyzes the reference document to produce a ReferenceSpec.
        """
        print("[NODE] Reference Agent: Inspecting reference document(s)...")
        if state.get("reference_pdf_paths") and len(state["reference_pdf_paths"]) > 0:
            try:
                # Use the new VisionAnalyzer to extract layout style from the uploaded PDFs/Images
                analyzer = VisionAnalyzer(state["reference_pdf_paths"], self.generator)
                state["reference_spec"] = analyzer.analyze()
                state["reference_images_b64"] = getattr(analyzer, "all_b64_images", [])
                
                # Extract text from ALL PDFs for factual context
                import pymupdf
                full_text = ""
                for path in state["reference_pdf_paths"]:
                    doc = pymupdf.open(path)
                    for page in doc:
                        full_text += page.get_text() + "\n"
                state["rag_context"] = full_text[:8000] # Cap at ~8k chars to respect free tier TPM limits
                
                print("[NODE] Reference Agent: Extracted Style and Factual Text successfully.")
            except Exception as e:
                print(f"[NODE] Reference Agent Error: {e}")
                state["reference_spec"] = None
        else:
            print("[NODE] Reference Agent: No reference document provided.")
            state["reference_spec"] = None
        return state

    def content_agent_node(self, state: AgentState) -> AgentState:
        """
        Since context is now dynamically pulled from the uploaded Reference PDF, 
        this node simply validates if we have enough context or bypasses.
        """
        print("[NODE] Content Agent: Checking provided context...")
        if not state.get("rag_context") and state.get("rag_topic_query"):
            # If the user asked for a topic but provided no PDF, we rely on Groq's internal GK.
            print(f"[NODE] Content Agent: No PDF provided. Delegating topic '{state['rag_topic_query']}' to Document Agent's general knowledge.")
        return state

    def planner_node(self, state: AgentState) -> AgentState:
        """
        Acts as Document Planner. Includes bounded A2A negotiation with Content Agent.
        """
        print("[NODE] Planner: Strategizing...")
        
        # A2A Negotiation Logic
        if state.get("rag_topic_query") and state.get("negotiation_count", 0) < 2:
            print("[NODE] Planner: Evaluating retrieved facts from Content Agent...")
            eval_sys = "You are the Document Planner. Your job is to verify if the provided context is sufficient to complete the user instruction. Reply exactly with 'YES' if it is sufficient. If insufficient, reply 'NO:' followed by a 1-sentence instruction of what specific facts are missing."
            eval_user = f"INSTRUCTION: {state['user_instruction']}\nCONTEXT:\n{state.get('rag_context', '')}"
            
            try:
                # We mock a simple bypass if mock is true
                if getattr(self.generator, "mock", False):
                    eval_result = "YES"
                else:
                    eval_result = self.generator.generate(eval_sys, eval_user).strip()
                    
                if eval_result.startswith("NO"):
                    feedback = eval_result.replace("NO:", "").strip()
                    print(f"[NODE] Planner -> Content Agent (A2A): Facts insufficient. {feedback}")
                    state["negotiation_feedback"] = feedback
                    state["negotiation_count"] += 1
                    return state # Will route back to content_agent
            except Exception as e:
                print(f"[NODE] Planner A2A Evaluation Error: {e}")
                
        print("[NODE] Planner: Facts verified. Proceeding to generation.")
        state["negotiation_feedback"] = None
        state["strategy"] = "Draft full document JSON in one pass."
        state["retry_count"] = 0
        state["validation_errors"] = None
        return state

    def writer_node(self, state: AgentState) -> AgentState:
        """
        Calls the LLM to generate the JSON, incorporating any validation errors if it's a retry.
        """
        print(f"[NODE] Writer: Generating JSON (Attempt {state['retry_count'] + 1})...")
        
        sys_prompt = build_system_prompt(state.get("reference_spec"), state.get("rag_context"), state.get("user_instruction", ""))
        prompt = f"{sys_prompt}\nNow, generate the JSON for the following instruction:\nUSER INSTRUCTION: {state['user_instruction']}\nOUTPUT JSON:"
        
        # If we failed previously, append the error to the prompt!
        if state.get("validation_errors") and state.get("raw_json"):
            prompt += f"\n\nYOUR PREVIOUS OUTPUT WAS INVALID.\n{state['raw_json']}\n\nERROR:\n{state['validation_errors']}\n\nPLEASE FIX THE JSON AND TRY AGAIN:\nOUTPUT JSON:"

        if getattr(self.generator, "mock", False):
            # Rig the mock to fail on attempt 1, succeed on attempt 2 (just like we did in Phase 7)
            if state["retry_count"] == 0:
                json_str = '{"document_type": "Quiz", "metadata": {}, "sections": [{"title": "Bad Quiz"}]}'
            else:
                json_str = '{"document_type": "Quiz", "metadata": {}, "page": {"size": "A4", "orientation": "portrait", "margins": {"top_mm": 20.0, "bottom_mm": 20.0, "left_mm": 20.0, "right_mm": 20.0}}, "theme": {"body": {"font_family": "Arial", "size_pt": 11.0, "bold": false, "italic": false, "underline": false, "color": {"hex_code": "#000000"}}, "h1": {"font_family": "Arial", "size_pt": 16.0, "bold": true, "italic": false, "underline": false, "color": {"hex_code": "#000000"}}, "h2": {"font_family": "Arial", "size_pt": 14.0, "bold": true, "italic": false, "underline": false, "color": {"hex_code": "#000000"}}, "h3": {"font_family": "Arial", "size_pt": 12.0, "bold": true, "italic": false, "underline": false, "color": {"hex_code": "#000000"}}, "answer_key": {"font_family": "Arial", "size_pt": 11.0, "bold": false, "italic": true, "underline": false, "color": {"hex_code": "#FF0000"}}}, "sections": [{"title": "Math Pop Quiz", "page_break_before": false, "elements": [{"element_type": "text", "content": "Instructions: Read the questions carefully and select the best answer.", "is_heading": false, "heading_level": 1, "alignment": "left", "style_override": null}, {"element_type": "text", "content": "1. What is 5 + 7?", "is_heading": false, "heading_level": 1, "alignment": "left", "style_override": null}, {"element_type": "list", "list_type": "lettered", "items": [{"content": "10"}, {"content": "11"}, {"content": "12"}, {"content": "13"}]}]}]}'
        else:
            try:
                # Call the abstracted generate method and pass the reference images so it can see the template!
                json_str = self.generator.generate(sys_prompt, prompt, images=state.get("reference_images_b64"))
                
                if "```json" in json_str:
                    json_str = json_str.split("```json")[1].split("```")[0].strip()
                elif "{" in json_str and "}" in json_str:
                    json_str = json_str[json_str.find("{"):json_str.rfind("}")+1]
            except Exception as e:
                json_str = "{}"
                safe_err = str(e).encode('ascii', 'replace').decode('ascii')
                print(f"Generation error: {safe_err}")
                
                # Surface the actual API error to the user and prevent pointless retries
                state["validation_errors"] = f"AI API Error: {safe_err}"
                state["retry_count"] = 99 # Force exit without retries
                
        state["raw_json"] = json_str
        state["retry_count"] += 1
        return state

    def critic_node(self, state: AgentState) -> AgentState:
        """
        Validates the JSON against the Pydantic schema.
        If valid, sets final_document. If invalid, sets validation_errors.
        """
        print("[NODE] Critic: Validating JSON...")
        
        # Bypass if we already hit a fatal error (like Rate Limit 413) upstream
        if state.get("retry_count", 0) >= 99:
            print(f"[NODE] Critic: Bypassing due to fatal error upstream: {state.get('validation_errors')}")
            return state
            
        try:
            data = json.loads(state["raw_json"])
            spec = DocumentSpec(**data)
            print("[NODE] Critic: PERFECT! Schema is valid.")
            state["final_document"] = spec
            state["validation_errors"] = None
        except json.JSONDecodeError as e:
            error_msg = f"JSONDecodeError: {str(e)}. Ensure the output is strictly valid JSON with no trailing commas or missing brackets."
            print(f"[NODE] Critic: FAIL (JSON Syntax) - {error_msg}")
            state["validation_errors"] = error_msg
        except ValidationError as e:
            error_msg = f"Pydantic ValidationError:\n{str(e)}\nEnsure all required fields are present and types are correct."
            print(f"[NODE] Critic: FAIL (Pydantic Schema) - {error_msg}")
            state["validation_errors"] = error_msg
        except Exception as e:
            error_msg = f"Unknown Error: {str(e)}"
            print(f"[NODE] Critic: FAIL (Unknown) - {error_msg}")
            state["validation_errors"] = error_msg
            
        return state

    def template_writer_node(self, state: AgentState) -> AgentState:
        """
        Generates a ReplacementMap for DOCX Template Injection.
        """
        print(f"[NODE] Template Writer: Generating ReplacementMap (Attempt {state['retry_count'] + 1})...")
        from docx import Document
        
        try:
            doc = Document(state["reference_docx_path"])
            template_text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
            # Add table text too for completeness
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            template_text += "\n" + cell.text.strip()
        except Exception as e:
            template_text = f"Error reading template: {str(e)}"
            
        sys_prompt = """You are an elite Template Injection AI. Your job is to output a strictly valid JSON ReplacementMap to transform an existing document template into a new document.
        
CRITICAL RULES:
1. You MUST replace EVERY SINGLE question, multiple-choice option, true/false statement, fill-in-the-blank, and answer key entry to match the new topic.
2. DO NOT leave old topic data (like old MCQ options) behind. If the question changes, its options MUST change.
3. You must create a SEPARATE replacement dictionary for each specific option if they are separated.
4. "old_text" MUST exactly match a substring in the TEMPLATE TEXT. Be precise.

EXAMPLE:
If the template has:
"i. What is the capital of France?"
"a. Berlin"
"b. Paris"
"c. Madrid"

And your new context requires a question about Science, you MUST output:
{
  "replacements": [
    {"old_text": "i. What is the capital of France?", "new_text": "i. What is the chemical symbol for Water?"},
    {"old_text": "a. Berlin", "new_text": "a. CO2"},
    {"old_text": "b. Paris", "new_text": "b. H2O"},
    {"old_text": "c. Madrid", "new_text": "c. O2"}
  ]
}
Notice how EVERY SINGLE option is replaced individually!
"""
        prompt = f"""
TEMPLATE TEXT:
{template_text}

USER INSTRUCTION:
{state['user_instruction']}

RAG CONTEXT (If applicable):
{state.get('rag_context', '')}

Output a strictly valid JSON matching this schema:
{{
  "replacements": [
    {{"old_text": "Exact text from template", "new_text": "New text based on context"}}
  ]
}}
"""
        if state.get("validation_errors") and state.get("raw_json"):
            prompt += f"\n\nYOUR PREVIOUS OUTPUT WAS INVALID.\n{state['raw_json']}\n\nERROR:\n{state['validation_errors']}\n\nPLEASE FIX THE JSON AND TRY AGAIN:\nOUTPUT JSON:"

        try:
            json_str = self.generator.generate(sys_prompt, prompt)
            if "```json" in json_str:
                json_str = json_str.split("```json")[1].split("```")[0].strip()
            elif "{" in json_str and "}" in json_str:
                json_str = json_str[json_str.find("{"):json_str.rfind("}")+1]
        except Exception as e:
            json_str = "{}"
            safe_err = str(e).encode('ascii', 'replace').decode('ascii')
            print(f"Generation error: {safe_err}")
            
            # Surface the actual API error to the user and prevent pointless retries
            state["validation_errors"] = f"AI API Error: {safe_err}"
            state["retry_count"] = 99 # Force exit without retries
            
        state["raw_json"] = json_str
        state["retry_count"] += 1
        return state

    def template_critic_node(self, state: AgentState) -> AgentState:
        """
        Validates the ReplacementMap JSON.
        """
        print("[NODE] Template Critic: Validating JSON...")
        if state.get("retry_count", 0) >= 99:
            return state
            
        try:
            data = json.loads(state["raw_json"])
            from schema.replacement_spec import ReplacementMap
            spec = ReplacementMap(**data)
            print("[NODE] Template Critic: PERFECT! Schema is valid.")
            state["replacement_map"] = spec
            state["validation_errors"] = None
        except Exception as e:
            error_msg = f"Validation Error: {str(e)}"
            print(f"[NODE] Template Critic: FAIL - {error_msg}")
            state["validation_errors"] = error_msg
            
        return state

import asyncio
from models.api.cloud_router import CloudRouter
from docx import Document

generator = CloudRouter(provider="openrouter", model="openai/gpt-4o-mini")

doc = Document("C:/Users/bhatt/OneDrive/Desktop/DocTree/G8_English_Literature_Worksheet.docx")
template_text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            if cell.text.strip():
                template_text += "\n" + cell.text.strip()

sys_prompt = """You are an elite Template Injection AI. Your job is to output a strictly valid JSON ReplacementMap to transform an existing document template into a new document.

CRITICAL RULES:
1. You MUST replace EVERY SINGLE question, multiple-choice option, true/false statement, fill-in-the-blank, and answer key entry to match the new topic.
2. DO NOT leave old topic data (like old MCQ options) behind. If the question changes, its options MUST change.
3. You must create a SEPARATE replacement dictionary for each specific option if they are separated.
4. "old_text" MUST exactly match a substring in the TEMPLATE TEXT. Be precise.
"""

prompt = f"""
TEMPLATE TEXT:
{template_text}

USER INSTRUCTION:
I have given you the G8_English_Literature_Worksheet.docx for refrence, your job is to take the how i taught my grandmother pdf and create all those types of questions, as given in our G8_English_Literature_Worksheet.docx for refrence, and put them in the G8_English_Literature_Worksheet.docx and give it back to me

Output a strictly valid JSON matching this schema:
{{
  "replacements": [
    {{"old_text": "Exact text from template", "new_text": "New text based on context"}}
  ]
}}
"""

res = generator.generate(sys_prompt, prompt)
print(res)
